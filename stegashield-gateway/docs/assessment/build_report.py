"""Build the assessment artefacts from reviewed Markdown. No application changes.
Requires python-docx and Pillow; generated diagrams are code-native illustrations.
"""
from pathlib import Path
import html
import re
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FIG = ROOT / 'figures'
FIG.mkdir(exist_ok=True)
INK = '#163b38'
GREEN = '#197965'
MUTED = '#58716d'
PALE = '#eef5f2'
GOLD = '#b57926'
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'

def font(size=25, bold=False):
    return ImageFont.truetype(BOLD if bold else FONT, size)

def canvas(title, subtitle, height=870):
    im = Image.new('RGB', (1600, height), 'white')
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1600, 10), fill=GREEN)
    d.text((50, 38), title, font=font(37, True), fill=INK)
    d.text((50, 94), subtitle, font=font(23), fill=MUTED)
    return im, d

def box(d, xy, title, lines, fill=PALE):
    x,y,w,h=xy
    d.rounded_rectangle((x,y,x+w,y+h), radius=17, fill=fill, outline='#bdd4cb', width=2)
    d.text((x+22,y+20), title, font=font(26,True), fill=INK)
    for i,line in enumerate(lines):
        d.text((x+22,y+64+i*34),line,font=font(22),fill=MUTED)

def arrow(d, points, label=None):
    d.line(points,fill=GREEN,width=4)
    x,y=points[-1]; px,py=points[-2]
    a=math.atan2(y-py,x-px)
    d.polygon([(x,y),(x-17*math.cos(a-.45),y-17*math.sin(a-.45)),(x-17*math.cos(a+.45),y-17*math.sin(a+.45))],fill=GREEN)
    if label:
        x,y=points[len(points)//2]
        d.text((x+8,y-32),label,font=font(19),fill=MUTED)

def save(im,name):
    im.save(FIG / (name+'.png'),dpi=(220,220))

im,d=canvas('01 / SYSTEM ARCHITECTURE','Implemented components and credential boundaries')
box(d,(55,180,360,180),'Browser portal',['React / TypeScript','Member and administrator','Tab-scoped session storage'])
box(d,(610,180,385,180),'FastAPI gateway',['JWT + trusted profile role','Permission and integrity checks','In-memory watermark adapters'])
box(d,(1140,180,405,180),'Supabase Auth',['Login and refresh','Signed access token / JWKS','Session claim when present'])
arrow(d,[(415,245),(610,245)])
arrow(d,[(995,245),(1140,245)])
arrow(d,[(225,180),(225,148),(1340,148),(1340,180)])
d.text((600,119),'Browser login / refresh',font=font(18),fill=MUTED)
box(d,(450,540,490,185),'PostgreSQL + RLS',['Caller-scoped catalogue and lookup','Trusted RPC: record before release','Private identity-to-token mapping'])
box(d,(1055,540,490,185),'Private original Storage',['Server credentials only','Original bytes unchanged','No public original links'])
arrow(d,[(725,360),(725,465),(695,465),(695,540)])
arrow(d,[(900,360),(900,445),(1300,445),(1300,540)])
d.text((60,785),'Trust boundary: browser never receives server credentials or unrestricted token mappings.',font=font(25,True),fill=GREEN)
save(im,'architecture')

im,d=canvas('02 / PROTECTED-COPY ISSUANCE','The same protected download operation powers preview and file download',820)
items=[('1. Authenticate',['Verify JWT and app role']),('2. Authorise',['Active document + grants']),('3. Validate original',['Check size / SHA-256']),('4. Generate copy',['Fresh UUID + HMAC']),('5. Commit event',['Recheck grants / session']),('6. Release bytes',['Only after valid event ID'])]
for i,(title,lines) in enumerate(items):
    row=i//3; col=i%3
    if row: col=2-col
    x=55+col*515;y=185+row*285
    box(d,(x,y,455,165),title,lines)
    if i in (0,1): arrow(d,[(x+455,y+83),(x+515,y+83)])
    if i==2: arrow(d,[(x+225,y+165),(x+225,y+285)])
    if i in (3,4): arrow(d,[(x,y+83),(x-60,y+83)])
d.text((65,715),'Failure at authorisation, integrity or event recording => no protected-file response.',font=font(26,True),fill=GOLD)
save(im,'download')

im,d=canvas('03 / FORENSIC DECISION PATH','A valid token alone is insufficient for a positive attribution',860)
left=[('Admin submits suspect file',['Append attempt-started audit']),('Validate + extract',['Recover one valid signed token']),('Resolve + compare',['Registered event + exact file hash'])]
for i,(title,lines) in enumerate(left):
    y=170+i*190;box(d,(80,y,620,140),title,lines)
    if i<2: arrow(d,[(390,y+140),(390,y+190)])
box(d,(960,170,550,150),'Validation / service error',['No attribution disclosed','An unfinished audit may remain'],fill='#fff5e8')
box(d,(960,395,550,150),'Inconclusive',['Missing / invalid / unknown token','Or issued-file hash differs'],fill='#fff5e8')
box(d,(960,630,550,155),'Verified copy match',['Append completed audit first','Return allowlisted event fields'])
arrow(d,[(700,240),(960,240)])
arrow(d,[(700,430),(820,430),(820,470),(960,470)])
arrow(d,[(700,620),(820,620),(820,705),(960,705)])
arrow(d,[(700,595),(870,595),(870,500),(960,500)])
d.text((720,565),'No exact match',font=font(18),fill=GOLD)
d.text((825,735),'Exact match',font=font(18),fill=GREEN)
d.text((80,780),'A match identifies an issued copy, not proof of who leaked it.',font=font(25,True),fill=GREEN)
save(im,'forensics')

im,d=canvas('04 / DATA RELATIONSHIPS','Logical overview; relationship labels show cardinality, not query direction',1110)
box(d,(590,170,420,125),'profiles',['id / role / display_name'])
box(d,(70,410,420,140),'sessions',['user_id -> profiles','optional event linkage'])
box(d,(590,410,420,140),'download_events',['user + document + session','unique token / protected hash'])
box(d,(1110,410,420,140),'documents',['private path / original hash','format / uploader'])
arrow(d,[(590,235),(280,235),(280,410)],'1 : many')
arrow(d,[(800,295),(800,410)],'1 : many')
arrow(d,[(1110,480),(1010,480)],'1 : many')
arrow(d,[(490,480),(590,480)])
box(d,(590,710,420,140),'forensic_events',['attempt_id / outcome','matched event + token'])
arrow(d,[(800,550),(800,710)],'1 : many')
box(d,(70,710,420,230),'Position access',['job_positions','position_memberships','document_position_permissions','Members: one current position'])
box(d,(1110,710,420,230),'Individual access',['document_permissions','document + user','expiry or no expiry','Union with position grants'])
# Access tables are grouped, not falsely shown as direct event foreign keys.
arrow(d,[(1320,710),(1320,550)])
d.text((65,1010),'Access groups link profiles/positions to documents; event mappings are admin-only.',font=font(25,True),fill=GREEN)
save(im,'data_model')

im,d=canvas('05 / AUTOMATED VERIFICATION','Recorded local test-case counts — not coverage, accuracy or project completion',740)
for i,(label,value) in enumerate([('Python backend',136),('PostgreSQL / pgTAP',70),('Frontend unit/component',12),('Browser scenarios',7)]):
    y=185+i*110
    d.text((55,y+13),label,font=font(27),fill=INK)
    d.rounded_rectangle((460,y,460+value*6.6,y+57),radius=7,fill=GREEN if i==0 else '#82b9a9')
    d.text((480+value*6.6,y+12),str(value),font=font(27,True),fill=INK)
d.text((55,655),'Backend rerun: 08 Sep 2026. Other counts: previously recorded local checkpoint.',font=font(23),fill=MUTED)
save(im,'test_counts')

im,d=canvas('06 / COMPLETION ROADMAP','Proposed dependency order; dates and individual assignments require group confirmation',780)
steps=[('A. Reconcile proposal',['Session vs download token','PDF carrier / photo scope']),('B. Close evidence gaps',['Session + IP in admin result','Hosted migration/demo record']),('C. Run research evaluation',['30–50-document corpus','Eight metrics + raw results']),('D. Finalise and defend',['Hardening / documentation','Verified logbooks + viva'])]
for i,(title,lines) in enumerate(steps):
    x=55+(i%2)*790;y=180+(i//2)*250
    box(d,(x,y,680,175),title,lines)
arrow(d,[(735,267),(845,267)])
arrow(d,[(1185,355),(1185,390),(395,390),(395,430)])
arrow(d,[(735,517),(845,517)])
d.text((55,695),'Progress review is a checkpoint, not a claim that research or deployment is complete.',font=font(25,True),fill=GREEN)
save(im,'roadmap')

source=(ROOT/'PROGRESS_REPORT.md').read_text(encoding='utf-8')
doc=Document()
section=doc.sections[0]
section.page_height=Inches(11.7);section.page_width=Inches(8.3)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.7)
styles=doc.styles
styles['Normal'].font.name='Calibri';styles['Normal'].font.size=Pt(10)
styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',34),('Heading 1',23),('Heading 2',16),('Heading 3',12)]:
    styles[name].font.name='Calibri';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string('197965')
styles['Heading 1'].paragraph_format.page_break_before=True
header=section.header.paragraphs[0]
header.text='GROUP 60  /  STEGASHIELD GATEWAY                                       PROGRESS REVIEW'
header.style='Caption'
footer=section.footer.paragraphs[0]
footer.text='08 September 2026  |  Review draft                                        '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

html_parts=[]; lines=source.splitlines(); i=0; cover=True
def inline(s):
    return re.sub(r'\*\*(.*?)\*\*',r'<strong>\1</strong>',html.escape(s))
while i<len(lines):
    line=lines[i]
    if not line.strip(): i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[c.strip() for c in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',c) for c in cells): rows.append(cells)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for n,row in enumerate(rows):
            cells=table.add_row().cells
            for cell,text in zip(cells,row):
                cell.text=text
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size=Pt(8);run.bold=n==0
            if n==0:
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        html_parts.append('<table><thead><tr>'+''.join('<th>'+inline(c)+'</th>' for c in rows[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inline(c)+'</td>' for c in row)+'</tr>' for row in rows[1:])+'</tbody></table>')
        continue
    match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
    if match:
        doc.add_picture(str(ROOT/match[2]),width=Inches(6.75))
        doc.add_paragraph(match[1],'Caption')
        html_parts.append(f'<figure><img src="{match[2]}"><figcaption>{html.escape(match[1])}</figcaption></figure>')
    elif line.startswith('# '):
        doc.add_paragraph(line[2:],'Title');html_parts.append('<header class="cover"><div class="kicker">INFORMATION SECURITY PROJECT 2026</div><h1>'+inline(line[2:])+'</h1>')
    elif line.startswith('## '):
        title=line[3:]
        if title=='Executive summary':
            doc.add_heading('Contents',level=1)
            for heading in re.findall(r'^## (.+)$',source,re.M)[1:]: doc.add_paragraph(heading)
            html_parts.append('</header><nav><h2>Contents</h2>'+''.join('<p>'+inline(h)+'</p>' for h in re.findall(r'^## (.+)$',source,re.M)[1:])+'</nav>')
        level=2 if cover else 1
        if title=='Executive summary': cover=False;level=1
        doc.add_heading(title,level=level)
        html_parts.append('<h2>'+inline(title)+'</h2>')
    elif line.startswith('### '):
        doc.add_heading(line[4:],level=2);html_parts.append('<h3>'+inline(line[4:])+'</h3>')
    else:
        doc.add_paragraph(line);html_parts.append('<p>'+inline(line)+'</p>')
    i+=1
doc.save(ROOT/'StegaShield_Group60_Progress_Report.docx')
css='''@page {size:A4; margin:18mm 17mm;} *{box-sizing:border-box} body{font:10.5pt/1.5 Arial,sans-serif;color:#183b38;max-width:185mm;margin:auto} h1{font-size:40pt;line-height:1.12;color:#197965} h2{font-size:21pt;color:#197965;margin-top:25px;break-after:avoid} h3{font-size:13pt;color:#197965;break-after:avoid} p{orphans:3;widows:3} .cover{padding-top:45px;border-top:10px solid #197965} .cover h2{break-before:auto;font-size:24pt}.kicker{font-size:10pt;letter-spacing:3px;color:#58716d} nav{break-before:page;break-after:page} nav p{margin:4px 0} body>h2{break-before:page} table{border-collapse:collapse;width:100%;font-size:8.5pt;margin:18px 0} th{background:#197965;color:white;text-align:left} th,td{padding:8px;border:1px solid #d7e5df;vertical-align:top;overflow-wrap:anywhere} tr:nth-child(even){background:#f0f6f3} tr{break-inside:avoid} thead{display:table-header-group} figure{margin:20px 0;break-inside:avoid} img{width:100%;height:auto} figcaption{font-size:9pt;color:#58716d;margin-top:8px} @media screen{body{padding:35px;background:white}html{background:#eaf1ee}}'''
(ROOT/'StegaShield_Group60_Progress_Report.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>StegaShield — Group 60 Progress Report</title><style>'+css+'</style><body>'+''.join(html_parts)+'</body></html>',encoding='utf-8')
print('Created Word, HTML and six high-resolution figures in',ROOT)
