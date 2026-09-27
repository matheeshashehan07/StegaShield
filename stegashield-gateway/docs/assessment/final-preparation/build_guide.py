"""Generate editable Word and printable HTML from the reviewed preparation guide."""
from pathlib import Path
import html
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
STEM = 'StegaShield_Final_Report_and_Viva_Guide'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.7)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
for level in range(1, 4):
    style = doc.styles[f'Heading {level}']
    style.font.color.rgb = RGBColor.from_string('176F60')
    style.paragraph_format.keep_with_next = True
section.header.paragraphs[0].text = 'STEGASHIELD  /  GROUP 60                                      FINAL ASSESSMENT PREPARATION'
section.header.paragraphs[0].style = 'Caption'
footer = section.footer.paragraphs[0]
footer.text = 'Preparation guide | Evidence must be verified                                      '
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def clean(s):
    return s.replace('**', '').replace('`', '')

parts = []
lines = (ROOT / 'FINAL_REPORT_AND_VIVA_GUIDE.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    i += 1
    if not line:
        continue
    if line.startswith('|'):
        rows = [line]
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip())
            i += 1
        cells = [[clean(c.strip()) for c in row.strip('|').split('|')] for row in rows
                 if not re.fullmatch(r'[|\s:\-]+', row)]
        table = doc.add_table(rows=0, cols=len(cells[0]))
        table.style = 'Light Shading Accent 1'
        parts.append('<table>')
        for n, row in enumerate(cells):
            target = table.add_row().cells
            parts.append('<tr>')
            for j, value in enumerate(row):
                target[j].text = value
                for p in target[j].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)
                        run.bold = n == 0
                tag = 'th' if n == 0 else 'td'
                parts.append(f'<{tag}>{html.escape(value)}</{tag}>')
            parts.append('</tr>')
            if n == 0:
                repeat = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
        parts.append('</table>')
        doc.add_paragraph()
    elif line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        value = clean(line[level:].strip())
        doc.add_heading(value, min(level, 3))
        parts.append(f'<h{level}>{html.escape(value)}</h{level}>')
    else:
        value = clean(line)
        style = 'Normal'
        if line.startswith('- '):
            value, style = clean(line[2:]), 'List Bullet'
        doc.add_paragraph(value, style)
        parts.append('<p>' + ('&#8226; ' if style == 'List Bullet' else '') + html.escape(value) + '</p>')

doc.core_properties.title = 'StegaShield: Final Report and Viva Preparation Guide'
doc.core_properties.subject = 'Final assessment expectations and evidence-based preparation'
doc.core_properties.author = 'StegaShield Group 60'
doc.save(ROOT / (STEM + '.docx'))
css = '''@page{size:A4;margin:19mm 17mm 20mm}body{font:10.5pt/1.42 Arial,sans-serif;color:#203733}
h1{font-size:26pt;border-bottom:4px solid #197965;padding-bottom:12px}h2{font-size:17pt;margin-top:24px}
h3{font-size:12pt}h1,h2,h3{color:#176f60;break-after:avoid}p{margin:0 0 9px}
table{width:100%;border-collapse:collapse;font-size:8.3pt;margin:12px 0}th{background:#e7f1ed;color:#164d41}
td,th{border:1px solid #c6d8d1;padding:6px;text-align:left;vertical-align:top}tr{break-inside:avoid}
thead{display:table-header-group}'''
(ROOT / (STEM + '.html')).write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>'
    + html.escape(doc.core_properties.title) + '</title><style>' + css + '</style><body>'
    + '\n'.join(parts) + '</body></html>', encoding='utf-8')
print('Generated Word and HTML preparation guide.')
