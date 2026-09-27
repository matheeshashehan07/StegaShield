from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from uuid import UUID
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph

from backend.app.core.errors import InvalidDocumentError
from backend.app.watermarking.base import WatermarkAdapter
from backend.app.watermarking.codec import SignedWatermarkCodec


@dataclass(frozen=True)
class ExtractedDocxWatermark:
    token: UUID
    valid_copies: int


class DocxWatermarkAdapter(WatermarkAdapter):
    MAX_REPETITIONS = 5

    def contains_watermark_frame(self, document: bytes) -> bool:
        """Reject protected originals, including mixed or invalid frames."""
        self._validate_package(document)
        try:
            doc = Document(BytesIO(document))
            return any(self.codec.START in p.text for p in self._all_paragraphs(doc))
        except Exception as exc:
            raise InvalidDocumentError("The DOCX could not be opened safely.") from exc

    def __init__(self, codec: SignedWatermarkCodec):
        self.codec = codec

    @staticmethod
    def _validate_package(data: bytes) -> None:
        try:
            with ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                members = archive.infolist()
                if len(members) > 2048 or len(names) != len(members) or sum(m.file_size for m in members) > 50 * 1024 * 1024:
                    raise InvalidDocumentError("The DOCX package exceeds safe expansion limits.")
                if any(m.flag_bits & 1 for m in members) or any("vbaproject" in n.lower() for n in names):
                    raise InvalidDocumentError("Encrypted or macro-enabled documents are unsupported.")
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise InvalidDocumentError("The file is not a valid DOCX document.")
                bad_member = archive.testzip()
                if bad_member is not None:
                    raise InvalidDocumentError("The DOCX package is corrupted.")
        except (BadZipFile, OSError):
            raise InvalidDocumentError("The file is not a valid DOCX document.") from None

    @staticmethod
    def _walk_table(table: Table):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested_table in cell.tables:
                    yield from DocxWatermarkAdapter._walk_table(nested_table)

    @classmethod
    def _body_paragraphs(cls, document: DocumentObject):
        yield from document.paragraphs
        for table in document.tables:
            yield from cls._walk_table(table)

    @classmethod
    def _all_paragraphs(cls, document: DocumentObject):
        yield from cls._body_paragraphs(document)
        seen_parts: set[str] = set()
        for section in document.sections:
            for container in (section.header, section.footer):
                part_name = str(container.part.partname)
                if part_name in seen_parts:
                    continue
                seen_parts.add(part_name)
                yield from container.paragraphs
                for table in container.tables:
                    yield from cls._walk_table(table)

    def embed(self, document: bytes, token: UUID) -> bytes:
        self._validate_package(document)
        try:
            doc = Document(BytesIO(document))
        except Exception as exc:
            raise InvalidDocumentError("The DOCX could not be opened safely.") from exc

        frame = self.codec.encode(token)
        candidates = [paragraph for paragraph in self._body_paragraphs(doc) if paragraph.text.strip()]
        if not candidates:
            raise InvalidDocumentError("The DOCX must contain at least one non-empty text paragraph.")

        # Spread repeated copies through the document to improve survival when
        # only part of a document is copied. Repetition does not expose identity.
        count = min(self.MAX_REPETITIONS, len(candidates))
        indexes = sorted({round(i * (len(candidates) - 1) / max(count - 1, 1)) for i in range(count)})
        for index in indexes:
            candidates[index].add_run(frame)

        output = BytesIO()
        doc.save(output)
        return output.getvalue()

    def extract(self, document: bytes) -> ExtractedDocxWatermark | None:
        self._validate_package(document)
        try:
            doc = Document(BytesIO(document))
        except Exception as exc:
            raise InvalidDocumentError("The DOCX could not be opened safely.") from exc

        tokens: list[UUID] = []
        for paragraph in self._all_paragraphs(doc):
            for frame in self.codec.find_frames(paragraph.text):
                decoded = self.codec.decode(frame)
                if decoded is not None:
                    tokens.append(decoded.token)

        if not tokens or len(set(tokens)) != 1:
            return None
        token, copies = Counter(tokens).most_common(1)[0]
        return ExtractedDocxWatermark(token=token, valid_copies=copies)
