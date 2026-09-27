from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from backend.app.core.auth import Principal, current_user
from backend.app.core.errors import InvalidDocumentError
from backend.app.main import app
from backend.app.services.document_service import DocumentWatermarkService
from backend.app.services.gateway import get_gateway, MIME


def make_docx():
    doc = Document()
    for n in range(6):
        doc.add_paragraph(f'Confidential paragraph {n}')
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class MemoryGateway:
    def __init__(self):
        self.data = make_docx()
        self.id = uuid4()
        self.metadata = dict(id=str(self.id), size_bytes=len(self.data), sha256=sha256(self.data).hexdigest(), format='docx')
        self.events = {}
        self.audits = []
        self.allowed = True
        self.fail_record = False
        self.fail_audit = False
        self.reads = 0
        self.uploads = []
        self.writes = []

    def store_original(self, path, data):
        self.uploads.append((path, data))

    def request(self, method, path, **kwargs):
        self.writes.append((method, path, kwargs))

    def document(self, document_id, user):
        if not self.allowed or document_id != self.id:
            raise HTTPException(404)
        return self.metadata

    def original(self, document):
        self.reads += 1
        return self.data

    def record_download(self, user, document_id, token, digest, ip, agent):
        if self.fail_record:
            raise HTTPException(503, 'Database unavailable')
        event = dict(id=str(uuid4()), user_id=str(user.id), document_id=str(document_id),
                     downloaded_at=datetime.now(timezone.utc).isoformat(), protected_sha256=digest)
        self.events[str(token)] = event
        return event['id']

    def lookup(self, token, user):
        return self.events.get(str(token))

    def audit(self, record):
        if self.fail_audit:
            raise HTTPException(503)
        self.audits.append(record.copy())


@pytest.fixture
def workflow():
    gateway = MemoryGateway()
    user = Principal(uuid4(), 'user', 'test')
    app.dependency_overrides[get_gateway] = lambda: gateway
    app.dependency_overrides[current_user] = lambda: user
    with TestClient(app) as client:
        yield client, gateway, user
    app.dependency_overrides.clear()


def admin():
    identity = Principal(uuid4(), 'admin', 'admin-test')
    app.dependency_overrides[current_user] = lambda: identity
    return identity


def extract(client, data, filename='suspect.docx'):
    return client.post('/api/v1/forensics/docx', files={'file': (filename, data, MIME)})


def test_authenticated_round_trip_and_unique_downloads(workflow):
    client, gateway, user = workflow
    original = gateway.data
    first = client.post(f'/api/v1/documents/{gateway.id}/download')
    second = client.post(f'/api/v1/documents/{gateway.id}/download')
    assert first.status_code == second.status_code == 200
    assert first.content != second.content
    assert len(gateway.events) == 2
    assert gateway.data == original
    assert first.headers['cache-control'] == 'no-store'
    assert 'x-watermark-token' not in first.headers
    examiner = admin()
    result = extract(client, first.content)
    assert result.status_code == 200
    assert result.json()['outcome'] == 'matched'
    assert result.json()['download']['user_id'] == str(user.id)
    assert result.json()['exact_copy'] is True
    assert gateway.audits[-1]['performed_by'] == str(examiner.id)
    assert gateway.audits[-1]['matched_download_id'] == result.json()['download']['id']


def test_denied_document_never_reads_storage(workflow):
    client, gateway, _ = workflow
    gateway.allowed = False
    assert client.post(f'/api/v1/documents/{gateway.id}/download').status_code == 404
    assert gateway.reads == 0
    assert not gateway.events


def test_database_failure_never_returns_protected_bytes(workflow):
    client, gateway, _ = workflow
    gateway.fail_record = True
    response = client.post(f'/api/v1/documents/{gateway.id}/download')
    assert response.status_code == 503
    assert not response.content.startswith(b'PK')
    assert 'content-disposition' not in response.headers


def test_storage_integrity_failure_prevents_record(workflow):
    client, gateway, _ = workflow
    gateway.data += b'changed'
    assert client.post(f'/api/v1/documents/{gateway.id}/download').status_code == 503
    assert not gateway.events


def test_ordinary_user_cannot_upload_grant_or_extract(workflow):
    client, gateway, _ = workflow
    assert extract(client, gateway.data).status_code == 403
    assert client.post('/api/v1/documents', data={'title': 'secret'}, files={'file': ('a.docx', gateway.data)}).status_code == 403
    assert client.post(f'/api/v1/documents/{gateway.id}/permissions', json={'user_id': str(uuid4())}).status_code == 403
    assert not gateway.audits


@pytest.mark.parametrize('kind', ['plain', 'unknown', 'modified', 'tampered', 'mixed'])
def test_inconclusive_never_discloses_identity(workflow, kind):
    client, gateway, _ = workflow
    data = client.post(f'/api/v1/documents/{gateway.id}/download').content
    service = DocumentWatermarkService.from_settings()
    if kind == 'plain':
        data = gateway.data
    elif kind == 'unknown':
        gateway.events.clear()
    elif kind == 'modified':
        doc = Document(BytesIO(data))
        doc.add_paragraph('Altered content')
        output = BytesIO(); doc.save(output); data = output.getvalue()
    elif kind == 'tampered':
        doc = Document(BytesIO(data))
        codec = service.docx_adapter.codec
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if codec.START in run.text:
                    run.text = run.text.replace(codec.START, '')
        output = BytesIO(); doc.save(output); data = output.getvalue()
    elif kind == 'mixed':
        data = service.embed_docx(data).document
    admin()
    result = extract(client, data)
    assert result.status_code == 200
    assert result.json()['outcome'] == 'inconclusive'
    assert 'download' not in result.json()
    assert gateway.audits[-1]['outcome'] == 'inconclusive'


def test_audit_failure_blocks_forensic_result(workflow):
    client, gateway, _ = workflow
    data = client.post(f'/api/v1/documents/{gateway.id}/download').content
    admin(); gateway.fail_audit = True
    assert extract(client, data).status_code == 503


@pytest.mark.parametrize('data,filename,status', [(b'broken', 'a.docx', 422), (b'', 'a.docx', 413), (b'pdf', 'a.pdf', 415)])
def test_invalid_attempts_audited(workflow, data, filename, status):
    client, gateway, _ = workflow
    admin()
    assert extract(client, data, filename).status_code == status
    assert gateway.audits[-1]['reason'] == 'invalid_document'


def test_zip_expansion_limit():
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '')
        archive.writestr('word/document.xml', b'x' * (50 * 1024 * 1024 + 1))
    with pytest.raises(InvalidDocumentError):
        DocumentWatermarkService.from_settings().extract_docx(output.getvalue())


def test_protected_original_rejected(workflow):
    client, gateway, _ = workflow
    gateway.data = DocumentWatermarkService.from_settings().embed_docx(gateway.data).document
    gateway.metadata.update(size_bytes=len(gateway.data), sha256=sha256(gateway.data).hexdigest())
    assert client.post(f'/api/v1/documents/{gateway.id}/download').status_code == 422
    assert not gateway.events


def test_admin_upload_derives_metadata_and_ignores_mime(workflow):
    client, gateway, _ = workflow
    identity = admin()
    result = client.post('/api/v1/documents', data={'title': 'Policy'},
        files={'file': ('../../malicious\".docx', gateway.data, 'text/plain')})
    assert result.status_code == 201
    row = gateway.writes[0][2]['json']
    assert row['uploaded_by'] == str(identity.id)
    assert row['mime_type'] == MIME
    assert row['sha256'] == sha256(gateway.data).hexdigest()
    assert gateway.uploads == [(f"{result.json()['id']}/original.docx", gateway.data)]
    assert row['original_filename'] == 'original.docx'


def test_invalid_upload_never_reaches_storage(workflow):
    client, gateway, _ = workflow
    admin()
    assert client.post('/api/v1/documents', data={'title': 'Policy'}, files={'file': ('a.docx', b'fake')}).status_code == 422
    assert gateway.uploads == []


def test_unacknowledged_record_never_returns_file(workflow, monkeypatch):
    client, gateway, _ = workflow
    monkeypatch.setattr(gateway, 'record_download', lambda *args: None)
    assert client.post(f'/api/v1/documents/{gateway.id}/download').status_code == 503


def test_client_identity_cannot_override_authenticated_user(workflow):
    client, gateway, identity = workflow
    response = client.post(f'/api/v1/documents/{gateway.id}/download', json={'user_id': str(uuid4()), 'token': str(uuid4()), 'role': 'admin'})
    assert response.status_code == 200
    assert next(iter(gateway.events.values()))['user_id'] == str(identity.id)


def test_lookup_failure_leaves_started_audit(workflow, monkeypatch):
    client, gateway, _ = workflow
    data = client.post(f'/api/v1/documents/{gateway.id}/download').content
    def fail(*args):
        raise HTTPException(503)
    monkeypatch.setattr(gateway, 'lookup', fail)
    admin()
    assert extract(client, data).status_code == 503
    assert gateway.audits[-1]['reason'] == 'attempt_started'
    assert gateway.audits[-1]['completed'] is False


def test_oversized_forensic_upload_is_audited(workflow, monkeypatch):
    from backend.app.api.v1 import documents
    from types import SimpleNamespace
    monkeypatch.setattr(documents, 'get_settings', lambda: SimpleNamespace(max_upload_size_bytes=1024))
    client, gateway, _ = workflow
    admin()
    assert extract(client, b'x' * 1025).status_code == 413
    assert gateway.audits[-1]['reason'] == 'invalid_document'


def test_conflicting_valid_tokens_have_no_winner():
    service = DocumentWatermarkService.from_settings()
    protected = service.embed_docx(make_docx()).document
    mixed = service.embed_docx(protected).document
    assert service.extract_docx(mixed) is None


def test_missing_authentication_on_new_routes(workflow):
    client, gateway, _ = workflow
    app.dependency_overrides.pop(current_user)
    assert client.get('/api/v1/documents').status_code == 401
    assert client.post(f'/api/v1/documents/{gateway.id}/download').status_code == 401
    assert extract(client, gateway.data).status_code == 401
