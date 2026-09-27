from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document

from backend.app.core.errors import InvalidDocumentError
from backend.app.watermarking.codec import SignedWatermarkCodec
from backend.app.watermarking.docx_adapter import DocxWatermarkAdapter


def make_docx(paragraph_count: int = 6) -> bytes:
    document = Document()
    for index in range(paragraph_count):
        document.add_paragraph(f"Confidential paragraph {index + 1}.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Confidential table value"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def visible_paragraphs(data: bytes) -> list[str]:
    document = Document(BytesIO(data))
    zero_width = {"\u200b", "\u200c", "\u200d", "\u2060", "\u2063"}
    return ["".join(c for c in paragraph.text if c not in zero_width) for paragraph in document.paragraphs]


def test_docx_round_trip_preserves_visible_paragraph_text() -> None:
    codec = SignedWatermarkCodec(b"a-secure-test-secret-with-at-least-32-bytes")
    adapter = DocxWatermarkAdapter(codec)
    original = make_docx()
    token = uuid4()
    protected = adapter.embed(original, token)
    extracted = adapter.extract(protected)
    assert extracted is not None
    assert extracted.token == token
    assert extracted.valid_copies == adapter.MAX_REPETITIONS
    assert visible_paragraphs(protected) == visible_paragraphs(original)


def test_non_watermarked_docx_returns_none() -> None:
    codec = SignedWatermarkCodec(b"a-secure-test-secret-with-at-least-32-bytes")
    adapter = DocxWatermarkAdapter(codec)
    assert adapter.extract(make_docx()) is None


def test_invalid_docx_is_rejected() -> None:
    codec = SignedWatermarkCodec(b"a-secure-test-secret-with-at-least-32-bytes")
    adapter = DocxWatermarkAdapter(codec)
    with pytest.raises(InvalidDocumentError):
        adapter.embed(b"not a docx", uuid4())
